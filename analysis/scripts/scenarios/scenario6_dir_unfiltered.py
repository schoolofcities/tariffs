"""
UNFILTERED Direct allocation for Scenario 6.

This is a deliberate variant of scenario_allocation.py's Direct effect.

Difference from the main script:
  - The main script restricts Direct to KNOWN_DIRECT_BS_CODES (the 2
    origin-of-shock codes from 'Delta FD Vectors': BS336110, BS336300).
    That's the NARROW definition -- "Direct = where the tariff struck."
  - This script does NOT filter. It auto-detects every BS code with a
    nonzero DIR value in DIR_INDIR_INDCD_SC6 (~192 codes), exactly the way
    Indirect/Induced are already handled. That's the BROAD definition --
    "Direct = the first-round (self/supplier) component of every industry's
    total loss, as the IO model itself decomposes it."

Rationale for the broad version (decided in chat):
  - Stays consistent with the model's own output instead of post-filtering it.
  - The largest DIR rows are, in fact, the auto production base -- motor
    vehicle parts, iron & steel mills, plastics, fabricated metal, forging/
    stamping, rubber, foundries, machine shops, auto-parts wholesalers, and
    automobile manufacturing itself -- so the layer is dominated by
    genuinely auto-related first-round activity.
  - DIR magnitudes are smaller than INDIR/INDCD, so this won't blow out the
    combined Total.

Caveat to confirm with DiFrancesco: a handful of large DIR rows are NOT
manufacturing (banking, credit unions, intangible-asset lessors, electric
power). They appear because they're big sectors with real first-round
adjustments, not because autos hit them directly. Label the layer as the
model's DIR decomposition, not as "industries the tariff hit."

Because broad Direct reaches ~192 industries across the whole economy (not
just the 2 auto codes), this MUST use trail_csd_full.csv -- trail_csd.csv
only covers the tariff-relevant NAICS and would silently zero out most of
the newly-included industries.

Output: scenario6_dir_csd_unfiltered.geojson  (+ .csv + _centroids.geojson)
Single Direct column: S6_DIR_Jobs (all negative = losses).
"""

import pandas as pd
import geopandas as gpd
import openpyxl
from pathlib import Path

# ============================================================
# PATHS -- must match scenario_allocation.py
# ============================================================
SCRIPT_DIR = Path(__file__).parent

MRIO_WORKBOOK = SCRIPT_DIR / '../raw/scenario summary - clsd MRIO - as of June 21 2026 - 2022 model.xlsx'
CONCORDANCE_WORKBOOK = SCRIPT_DIR / '../intermediate/NAICS 2022v1 to IOIC 2022 concordance.xlsx'
TRAIL_CSD_FULL_CSV = SCRIPT_DIR / '../intermediate/trail_csd_full.csv'  # REQUIRED for broad Direct
CSD_SHAPEFILE = SCRIPT_DIR / '../../data/census/lcsd000b21a_e/lcsd000b21a_e.shp'
OUTPUT_DIR = SCRIPT_DIR / '../outputs'

SCENARIO_NUM = 6
EFFECT = 'DIR'
EFFECT_COLUMN_INDEX = {'DIR': 4, 'INDIR': 5, 'INDCD': 6}
OUTPUT_COL = f'S{SCENARIO_NUM}_DIR_Jobs'
OUTPUT_NAME = f'scenario{SCENARIO_NUM}_dir_csd_unfiltered'

MRIO_TO_NOTEBOOK_PROVINCE = {
    'NL': 'NL', 'PEI': 'PEI', 'NS': 'NS', 'NB': 'NB', 'QC': 'QC',
    'ON': 'ON', 'MB': 'MB', 'SK': 'SK', 'AB': 'AL', 'BC': 'BC',
    'YT': 'YK', 'NT': 'NWT', 'NU': 'NU'
}

PROVINCE_CODE = {
    10: 'NL', 11: 'PEI', 12: 'NS', 13: 'NB', 24: 'QC', 35: 'ON',
    46: 'MB', 47: 'SK', 48: 'AL', 59: 'BC', 60: 'YK', 61: 'NWT', 62: 'NU'
}


def load_bs_to_naics(concordance_path):
    wb = openpyxl.load_workbook(concordance_path)
    ws = wb['English']
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    lookup = {}
    for r in rows:
        naics, naics_title, bs_code, bs_title = r[0], r[1], r[2], r[3]
        if bs_code is None or naics is None:
            continue
        lookup.setdefault(bs_code, []).append(str(naics))
    return lookup


BS_TO_NAICS = load_bs_to_naics(CONCORDANCE_WORKBOOK)


# ============================================================
# STEP 1: Extract provincial DIR job loss, per BS code (unchanged logic)
# ============================================================
def extract_provincial_effect_jobs(workbook_path, scenario_num, effect):
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    dir_sheet_name = f'DIR_INDIR_INDCD_SC{scenario_num}'
    effect_col = EFFECT_COLUMN_INDEX[effect]

    delta_x_rows = list(wb['Delta X - D Level'].iter_rows(min_row=4, values_only=True))
    delta_x_pb = []
    current_province = None
    for r in delta_x_rows:
        if r[0] is not None:
            current_province = r[0]
        delta_x_pb.append((current_province, r[3]))

    effect_rows = list(wb[dir_sheet_name].iter_rows(min_row=2, values_only=True))
    effect_province = []
    effect_values = []
    current_province = None
    for r in effect_rows:
        if r[1] is not None:
            current_province = r[1]
        effect_province.append(current_province)
        effect_values.append(r[effect_col])

    mult_rows = list(wb['Jobs Multipliers D'].iter_rows(min_row=2, values_only=True))
    multiplier_lookup = {}
    current_province = None
    for r in mult_rows:
        if r[1] is not None:
            current_province = r[1]
        multiplier_lookup[(current_province, r[2])] = r[4]

    def split_into_blocks(province_list, payload_list):
        blocks = []
        for i, prov in enumerate(province_list):
            if prov in (None, 'Canada'):
                continue
            item = payload_list[i]
            if blocks and blocks[-1][0] == prov:
                blocks[-1][1].append(item)
            else:
                blocks.append((prov, [item]))
        return blocks

    delta_x_blocks = split_into_blocks([p for p, _ in delta_x_pb], [bs for _, bs in delta_x_pb])
    effect_blocks = split_into_blocks(effect_province, effect_values)

    delta_x_by_prov = dict(delta_x_blocks)
    effect_by_prov = dict(effect_blocks)

    mismatches = {
        prov: (len(delta_x_by_prov[prov]), len(effect_by_prov.get(prov, [])))
        for prov in delta_x_by_prov
        if len(delta_x_by_prov[prov]) != len(effect_by_prov.get(prov, []))
    }
    if mismatches:
        print(f"⚠️  [{effect}] Per-province block size mismatch: {mismatches}")
        print("    Truncating mismatched provinces to the shorter length -- verify if unexpected.")

    records = []
    for prov, bs_codes in delta_x_blocks:
        effect_vals_for_prov = effect_by_prov.get(prov, [])
        n = min(len(bs_codes), len(effect_vals_for_prov))
        for bs_code, dollars in zip(bs_codes[:n], effect_vals_for_prov[:n]):
            mult = multiplier_lookup.get((prov, bs_code))
            if mult is None:
                continue
            jobs = (dollars or 0) / 1_000_000 * mult
            records.append({'Province': prov, 'BS_Code': bs_code, 'Jobs': jobs})

    df = pd.DataFrame(records)

    # NO KNOWN_DIRECT_BS_CODES cross-check here -- that's the whole point of
    # the unfiltered variant. We keep every nonzero DIR row.
    provincial_totals = df.groupby('Province')['Jobs'].sum()
    print(f"   [{effect}] National total: {provincial_totals.sum():,.1f} jobs")
    print(f"   [{effect}] Nonzero BS codes (unfiltered): "
          f"{df[df['Jobs'].abs() > 1e-6]['BS_Code'].nunique()}")

    return df


def get_relevant_bs_codes(df, threshold=1e-6):
    nonzero = df[df['Jobs'].abs() > threshold]
    return sorted(nonzero['BS_Code'].unique())


# ============================================================
# STEP 2: Load CSD employment (unchanged)
# ============================================================
def load_csd_employment(trail_csd_path, bs_codes):
    df = pd.read_csv(trail_csd_path, dtype={'CSDDGUID': str})

    needed_naics = sorted(set(n for bs in bs_codes for n in BS_TO_NAICS.get(bs, [])))
    missing = [c for c in needed_naics if c not in df.columns]
    if missing:
        print(f"⚠️  {len(missing)} NAICS columns not found in {trail_csd_path.name} "
              f"(first 10: {missing[:10]}). Those industries contribute zero. "
              "If this list is large, confirm you're pointing at trail_csd_full.csv.")

    available_naics = [c for c in needed_naics if c in df.columns]

    df['CSDUID'] = df['CSDDGUID'].str[-7:]
    df['PRUID'] = df['CSDUID'].str[:2].astype(int)
    df['Province'] = df['PRUID'].map(PROVINCE_CODE)

    return df[['CSDDGUID', 'Province'] + available_naics]


# ============================================================
# STEP 3: Allocate (unchanged)
# ============================================================
def allocate_to_csd(csd_employment_df, provincial_effect_jobs_df, bs_codes, output_col):
    result = csd_employment_df[['CSDDGUID', 'Province']].copy()
    result[output_col] = 0.0

    jobs_df = provincial_effect_jobs_df.copy()
    jobs_df['Province'] = jobs_df['Province'].map(MRIO_TO_NOTEBOOK_PROVINCE)

    skipped_no_concordance = []
    skipped_no_naics_data = []

    for bs_code in bs_codes:
        naics_cols = BS_TO_NAICS.get(bs_code)
        if not naics_cols:
            skipped_no_concordance.append(bs_code)
            continue

        available_cols = [c for c in naics_cols if c in csd_employment_df.columns]
        if not available_cols:
            skipped_no_naics_data.append(bs_code)
            continue

        csd_bs_employment = csd_employment_df[available_cols].sum(axis=1)
        prov_bs_totals = csd_bs_employment.groupby(csd_employment_df['Province']).sum().to_dict()

        weight = csd_employment_df['Province'].map(
            lambda p, totals=prov_bs_totals: totals.get(p, 0)
        )
        weight = csd_bs_employment / weight.replace(0, pd.NA)
        weight = weight.fillna(0.0)

        prov_loss_for_bs = jobs_df[jobs_df['BS_Code'] == bs_code].set_index('Province')['Jobs'].to_dict()
        provincial_loss = csd_employment_df['Province'].map(prov_loss_for_bs).fillna(0.0)

        result[output_col] += weight * provincial_loss

    if skipped_no_concordance:
        print(f"⚠️  {len(skipped_no_concordance)} BS codes had no concordance entry, skipped "
              f"(sample: {skipped_no_concordance[:5]})")
    if skipped_no_naics_data:
        print(f"⚠️  {len(skipped_no_naics_data)} BS codes had no matching NAICS columns in "
              f"the employment file, skipped (sample: {skipped_no_naics_data[:5]})")

    return result


# ============================================================
# STEP 4: Save (unchanged -- polygons + centroids)
# ============================================================
def save_outputs(result_df, csd_shapefile_path, output_name):
    csd = gpd.read_file(csd_shapefile_path)
    csd['CSDDGUID'] = csd['DGUID']
    csd = csd[['CSDDGUID', 'geometry']].copy()

    csd = csd[csd.geometry.notna() & ~csd.geometry.is_empty].copy()
    invalid = ~csd.geometry.is_valid
    if invalid.any():
        print(f"   Repairing {invalid.sum()} invalid CSD geometries via buffer(0)...")
        csd.loc[invalid, 'geometry'] = csd.loc[invalid, 'geometry'].buffer(0)
        csd = csd[~csd.geometry.is_empty].copy()

    merged = csd.merge(result_df, on='CSDDGUID', how='left')
    merged = gpd.GeoDataFrame(merged, geometry='geometry', crs=csd.crs)
    merged = merged.to_crs('EPSG:4326')

    geojson_path = OUTPUT_DIR / 'geojson' / f'{output_name}.geojson'
    csv_path = OUTPUT_DIR / 'csv' / f'{output_name}.csv'
    centroid_path = OUTPUT_DIR / 'geojson' / f'{output_name}_centroids.geojson'
    geojson_path.parent.mkdir(parents=True, exist_ok=True)
    csv_path.parent.mkdir(parents=True, exist_ok=True)

    merged.to_file(geojson_path, driver='GeoJSON')
    merged.drop(columns='geometry').to_csv(csv_path, index=False)

    centroids = merged.to_crs('EPSG:3347').copy()
    centroids['geometry'] = centroids.geometry.representative_point()
    centroids = centroids.to_crs('EPSG:4326')
    centroids.to_file(centroid_path, driver='GeoJSON')

    print(f"✅ Saved {geojson_path}")
    print(f"✅ Saved {csv_path}")
    print(f"✅ Saved {centroid_path}  (bubbles / Count mode)")


# ============================================================
# MAIN
# ============================================================
if __name__ == '__main__':
    print(f"=== Scenario {SCENARIO_NUM}, effect={EFFECT} (UNFILTERED Direct) ===")

    print("STEP 1: Extracting provincial DIR job loss (no origin-code filter)...")
    provincial_jobs = extract_provincial_effect_jobs(MRIO_WORKBOOK, SCENARIO_NUM, EFFECT)

    # THE key difference: auto-detect all nonzero DIR codes, exactly like
    # Indirect/Induced -- no KNOWN_DIRECT_BS_CODES restriction.
    bs_codes = get_relevant_bs_codes(provincial_jobs)
    print(f"   Using {len(bs_codes)} BS codes (unfiltered -- expect ~190)")

    print("STEP 2: Loading CSD-level employment (trail_csd_full.csv REQUIRED)...")
    csd_employment = load_csd_employment(TRAIL_CSD_FULL_CSV, bs_codes)

    print("STEP 3: Allocating to CSDs...")
    result = allocate_to_csd(csd_employment, provincial_jobs, bs_codes, OUTPUT_COL)
    print(f"   Total allocated: {result[OUTPUT_COL].sum():,.1f} jobs "
          f"(should ≈ provincial total: {provincial_jobs['Jobs'].sum():,.1f})")

    print("STEP 4: Saving...")
    save_outputs(result, CSD_SHAPEFILE, OUTPUT_NAME)
