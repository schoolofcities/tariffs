"""
CSD job-loss allocation for any scenario, any effect (Direct/Indirect/Induced).

Extends the original Scenario 6 Direct-only script:
  - Direct effect: reads DIR column, and (per scenario) is expected to hit
    only a handful of BS codes -- we cross-check against the known shocked
    codes rather than trusting Sheet1's text description blindly.
  - Indirect / Induced effects: read INDIR/INDCD columns respectively, and
    ripple across nearly all ~220 BS codes -- so instead of a hardcoded BS
    code list, we auto-detect every BS code with a nonzero value for that
    effect under this scenario, directly from the workbook.

Requires trail_csd_full.csv (see build_trail_csd_full.py) for Indirect/
Induced, since those ripple into industries outside the tariff-relevant
NAICS scope that trail_csd.csv is limited to. Direct effect still works
fine with the original trail_csd.csv if you don't need Indirect/Induced.
"""

import pandas as pd
import geopandas as gpd
import openpyxl
from pathlib import Path

# ============================================================
# PATHS -- adjust to match your project structure
# ============================================================
SCRIPT_DIR = Path(__file__).parent

MRIO_WORKBOOK = SCRIPT_DIR / '../raw/scenario summary - clsd MRIO - as of June 21 2026 - 2022 model.xlsx'
CONCORDANCE_WORKBOOK = SCRIPT_DIR / '../intermediate/NAICS 2022v1 to IOIC 2022 concordance.xlsx'
TRAIL_CSD_CSV = SCRIPT_DIR / '../intermediate/trail_csd.csv'            # Direct only (tariff-relevant NAICS)
TRAIL_CSD_FULL_CSV = SCRIPT_DIR / '../intermediate/trail_csd_full.csv'  # Indirect/Induced (full economy)
CSD_SHAPEFILE = SCRIPT_DIR / '../../data/census/lcsd000b21a_e/lcsd000b21a_e.shp'
OUTPUT_DIR = SCRIPT_DIR / '../outputs'

# Province label translation: MRIO workbook uses AB/YT/NT,
# your notebook's province_code dict uses AL/YK/NWT for the same provinces.
MRIO_TO_NOTEBOOK_PROVINCE = {
    'NL': 'NL', 'PEI': 'PEI', 'NS': 'NS', 'NB': 'NB', 'QC': 'QC',
    'ON': 'ON', 'MB': 'MB', 'SK': 'SK', 'AB': 'AL', 'BC': 'BC',
    'YT': 'YK', 'NT': 'NWT', 'NU': 'NU'
}

PROVINCE_CODE = {
    10: 'NL', 11: 'PEI', 12: 'NS', 13: 'NB', 24: 'QC', 35: 'ON',
    46: 'MB', 47: 'SK', 48: 'AL', 59: 'BC', 60: 'YK', 61: 'NWT', 62: 'NU'
}

# Which BS codes carry a real DIRECT shock, per scenario -- verified against
# the workbook's raw shock data ('Delta FD Vectors'), not just Sheet1's text
# description (which is sometimes wrong -- see Scenario 6's BS336120/
# BS336200 false positives, confirmed earlier).
KNOWN_DIRECT_BS_CODES = {
    6: ['BS336110', 'BS336300'],
    # Add other scenarios here once verified against 'Delta FD Vectors' the
    # same way -- don't assume Sheet1's text description is correct as-is.
}

EFFECT_COLUMN_INDEX = {'DIR': 4, 'INDIR': 5, 'INDCD': 6}


def load_bs_to_naics(concordance_path):
    """Real BS_Code -> [NAICS codes] lookup from the official StatCan concordance."""
    wb = openpyxl.load_workbook(concordance_path)
    ws = wb['English']
    rows = list(ws.iter_rows(min_row=3, values_only=True))

    lookup = {}
    for r in rows:
        naics, naics_title, bs_code, bs_title = r[0], r[1], r[2], r[3]
        if bs_code is None or naics is None:
            continue
        lookup.setdefault(bs_code, []).append(str(naics))
    return lookup


BS_TO_NAICS = load_bs_to_naics(CONCORDANCE_WORKBOOK)


# ============================================================
# STEP 1: Extract provincial job loss, per BS code, for a given effect
# ============================================================
def extract_provincial_effect_jobs(workbook_path, scenario_num, effect):
    """effect: one of 'DIR', 'INDIR', 'INDCD'"""
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    dir_sheet_name = f'DIR_INDIR_INDCD_SC{scenario_num}'
    effect_col = EFFECT_COLUMN_INDEX[effect]

    # ---- Delta X - D Level: (Province, BS_Code) per row ----
    delta_x_rows = list(wb['Delta X - D Level'].iter_rows(min_row=4, values_only=True))
    delta_x_pb = []
    current_province = None
    for r in delta_x_rows:
        if r[0] is not None:
            current_province = r[0]
        delta_x_pb.append((current_province, r[3]))

    # ---- DIR_INDIR_INDCD_SC{n}: has its OWN Province column ----
    effect_rows = list(wb[dir_sheet_name].iter_rows(min_row=2, values_only=True))
    effect_province = []
    effect_values = []
    current_province = None
    for r in effect_rows:
        if r[1] is not None:
            current_province = r[1]
        effect_province.append(current_province)
        effect_values.append(r[effect_col])

    # ---- Jobs Multipliers D: keyed lookup, no position dependency ----
    mult_rows = list(wb['Jobs Multipliers D'].iter_rows(min_row=2, values_only=True))
    multiplier_lookup = {}
    current_province = None
    for r in mult_rows:
        if r[1] is not None:
            current_province = r[1]
        multiplier_lookup[(current_province, r[2])] = r[4]

    # ---- Split into per-province blocks and validate sizes match ----
    def split_into_blocks(province_list, payload_list):
        blocks = []
        for i, prov in enumerate(province_list):
            if prov in (None, 'Canada'):
                continue
            item = payload_list[i]
            if blocks and blocks[-1][0] == prov:
                blocks[-1][1].append(item)
            else:
                blocks.append((prov, [item]))
        return blocks

    delta_x_blocks = split_into_blocks([p for p, _ in delta_x_pb], [bs for _, bs in delta_x_pb])
    effect_blocks = split_into_blocks(effect_province, effect_values)

    delta_x_by_prov = dict(delta_x_blocks)
    effect_by_prov = dict(effect_blocks)

    mismatches = {
        prov: (len(delta_x_by_prov[prov]), len(effect_by_prov.get(prov, [])))
        for prov in delta_x_by_prov
        if len(delta_x_by_prov[prov]) != len(effect_by_prov.get(prov, []))
    }
    if mismatches:
        print(f"⚠️  [{effect}] Per-province block size mismatch: {mismatches}")
        print("    Truncating mismatched provinces to the shorter length -- verify if unexpected.")

    records = []
    for prov, bs_codes in delta_x_blocks:
        effect_vals_for_prov = effect_by_prov.get(prov, [])
        n = min(len(bs_codes), len(effect_vals_for_prov))
        for bs_code, dollars in zip(bs_codes[:n], effect_vals_for_prov[:n]):
            mult = multiplier_lookup.get((prov, bs_code))
            if mult is None:
                continue
            jobs = (dollars or 0) / 1_000_000 * mult
            records.append({'Province': prov, 'BS_Code': bs_code, 'Jobs': jobs})

    df = pd.DataFrame(records)

    if effect == 'DIR' and scenario_num in KNOWN_DIRECT_BS_CODES:
        nonzero = df[df['Jobs'].abs() > 1e-6]
        unexpected = set(nonzero['BS_Code'].unique()) - set(KNOWN_DIRECT_BS_CODES[scenario_num])
        if unexpected:
            print(f"⚠️  WARNING: unexpected BS codes with nonzero Direct jobs: {unexpected}")
        else:
            print(f"✅ Confirmed: only {KNOWN_DIRECT_BS_CODES[scenario_num]} carry nonzero Direct jobs.")

    provincial_totals = df.groupby('Province')['Jobs'].sum()
    print(f"   [{effect}] National total: {provincial_totals.sum():,.1f} jobs")
    print(f"   [{effect}] Relevant BS codes: {df[df['Jobs'].abs() > 1e-6]['BS_Code'].nunique()}")

    return df


def get_relevant_bs_codes(provincial_effect_jobs_df, threshold=1e-6):
    """Auto-detect which BS codes actually carry a nonzero value for this
    effect -- for Indirect/Induced this will be most of the ~220 codes."""
    nonzero = provincial_effect_jobs_df[provincial_effect_jobs_df['Jobs'].abs() > threshold]
    return sorted(nonzero['BS_Code'].unique())


# ============================================================
# STEP 2: Load CSD-level raw employment for whichever NAICS codes are needed
# ============================================================
def load_csd_employment(trail_csd_path, bs_codes):
    df = pd.read_csv(trail_csd_path, dtype={'CSDDGUID': str})

    needed_naics = sorted(set(n for bs in bs_codes for n in BS_TO_NAICS.get(bs, [])))
    missing = [c for c in needed_naics if c not in df.columns]
    if missing:
        print(f"⚠️  {len(missing)} NAICS columns not found in {trail_csd_path.name} "
              f"(first 10: {missing[:10]}). These BS codes' contributions will be "
              "treated as zero -- if this is Indirect/Induced, you likely need "
              "trail_csd_full.csv instead of trail_csd.csv.")

    available_naics = [c for c in needed_naics if c in df.columns]

    df['CSDUID'] = df['CSDDGUID'].str[-7:]
    df['PRUID'] = df['CSDUID'].str[:2].astype(int)
    df['Province'] = df['PRUID'].map(PROVINCE_CODE)

    return df[['CSDDGUID', 'Province'] + available_naics]


# ============================================================
# STEP 3: Allocate provincial job loss to CSDs, per BS code, then sum
# ============================================================
def allocate_to_csd(csd_employment_df, provincial_effect_jobs_df, bs_codes, output_col):
    result = csd_employment_df[['CSDDGUID', 'Province']].copy()
    result[output_col] = 0.0

    jobs_df = provincial_effect_jobs_df.copy()
    jobs_df['Province'] = jobs_df['Province'].map(MRIO_TO_NOTEBOOK_PROVINCE)

    skipped_no_concordance = []
    skipped_no_naics_data = []

    for bs_code in bs_codes:
        naics_cols = BS_TO_NAICS.get(bs_code)
        if not naics_cols:
            skipped_no_concordance.append(bs_code)
            continue

        available_cols = [c for c in naics_cols if c in csd_employment_df.columns]
        if not available_cols:
            skipped_no_naics_data.append(bs_code)
            continue

        csd_bs_employment = csd_employment_df[available_cols].sum(axis=1)
        prov_bs_totals = csd_bs_employment.groupby(csd_employment_df['Province']).sum().to_dict()

        weight = csd_employment_df['Province'].map(
            lambda p, totals=prov_bs_totals: totals.get(p, 0)
        )
        weight = csd_bs_employment / weight.replace(0, pd.NA)
        weight = weight.fillna(0.0)

        prov_loss_for_bs = jobs_df[jobs_df['BS_Code'] == bs_code].set_index('Province')['Jobs'].to_dict()
        provincial_loss = csd_employment_df['Province'].map(prov_loss_for_bs).fillna(0.0)

        result[output_col] += weight * provincial_loss

    if skipped_no_concordance:
        print(f"⚠️  {len(skipped_no_concordance)} BS codes had no concordance entry, skipped "
              f"(sample: {skipped_no_concordance[:5]})")
    if skipped_no_naics_data:
        print(f"⚠️  {len(skipped_no_naics_data)} BS codes had no matching NAICS columns in "
              f"the employment file, skipped (sample: {skipped_no_naics_data[:5]})")

    return result


# ============================================================
# STEP 4: Save outputs
# ============================================================
def save_outputs(result_df, csd_shapefile_path, output_name):
    csd = gpd.read_file(csd_shapefile_path)
    csd['CSDDGUID'] = csd['DGUID']
    csd = csd[['CSDDGUID', 'geometry']].copy()

    # Drop null/empty geometries and repair invalid ones (self-intersections
    # etc.) BEFORE any reprojection -- shapely's transform throws an opaque
    # "bad allocation" GEOSException on malformed geometry, which is what
    # crashed the reproject step here.
    csd = csd[csd.geometry.notna() & ~csd.geometry.is_empty].copy()
    invalid = ~csd.geometry.is_valid
    if invalid.any():
        print(f"   Repairing {invalid.sum()} invalid CSD geometries via buffer(0)...")
        csd.loc[invalid, 'geometry'] = csd.loc[invalid, 'geometry'].buffer(0)
        # buffer(0) can still leave empties for degenerate slivers -- drop those
        csd = csd[~csd.geometry.is_empty].copy()

    merged = csd.merge(result_df, on='CSDDGUID', how='left')
    merged = gpd.GeoDataFrame(merged, geometry='geometry', crs=csd.crs)
    merged = merged.to_crs('EPSG:4326')

    geojson_path = OUTPUT_DIR / 'geojson' / f'{output_name}.geojson'
    csv_path = OUTPUT_DIR / 'csv' / f'{output_name}.csv'
    centroid_path = OUTPUT_DIR / 'geojson' / f'{output_name}_centroids.geojson'
    geojson_path.parent.mkdir(parents=True, exist_ok=True)
    csv_path.parent.mkdir(parents=True, exist_ok=True)

    # Polygon output (for choropleth / Percent mode)
    merged.to_file(geojson_path, driver='GeoJSON')
    merged.drop(columns='geometry').to_csv(csv_path, index=False)

    # Centroid output (for proportional bubbles / Count mode).
    # representative_point() guarantees a point INSIDE each polygon --
    # matches the notebook's method and avoids MapLibre placing circles on
    # polygon vertices/edges for concave or multipart CSD shapes.
    # Projected CRS first so the geometric operation is valid, then back to 4326.
    centroids = merged.to_crs('EPSG:3347').copy()
    centroids['geometry'] = centroids.geometry.representative_point()
    centroids = centroids.to_crs('EPSG:4326')
    centroids.to_file(centroid_path, driver='GeoJSON')

    print(f"✅ Saved {geojson_path}")
    print(f"✅ Saved {csv_path}")
    print(f"✅ Saved {centroid_path}  (bubbles / Count mode)")


# ============================================================
# ORCHESTRATION: run all three effects for one scenario, merge into
# ONE output (shared geometry, one column per effect + a Total column)
# ============================================================
def run_scenario_all_effects(scenario_num, effect_to_trail_csd):
    """effect_to_trail_csd: e.g. {'DIR': TRAIL_CSD_CSV,
    'INDIR': TRAIL_CSD_FULL_CSV, 'INDCD': TRAIL_CSD_FULL_CSV}"""
    effect_cols = []
    merged = None

    for effect, trail_csd_path in effect_to_trail_csd.items():
        output_col = f'S{scenario_num}_{effect}_Jobs'
        effect_cols.append(output_col)

        print(f"\n=== Scenario {scenario_num}, effect={effect} ===")
        print("STEP 1: Extracting provincial job loss...")
        provincial_jobs = extract_provincial_effect_jobs(MRIO_WORKBOOK, scenario_num, effect)

        bs_codes = (
            KNOWN_DIRECT_BS_CODES[scenario_num] if effect == 'DIR' and scenario_num in KNOWN_DIRECT_BS_CODES
            else get_relevant_bs_codes(provincial_jobs)
        )
        print(f"   Using {len(bs_codes)} BS codes")

        print("STEP 2: Loading CSD-level employment...")
        csd_employment = load_csd_employment(trail_csd_path, bs_codes)

        print("STEP 3: Allocating to CSDs...")
        result = allocate_to_csd(csd_employment, provincial_jobs, bs_codes, output_col)
        print(f"   Total allocated: {result[output_col].sum():,.1f} jobs "
              f"(should ≈ provincial total: {provincial_jobs['Jobs'].sum():,.1f})")

        # Keep only CSDDGUID + this effect's column for merging -- geometry
        # gets attached once, at the very end, in save_outputs.
        effect_only = result[['CSDDGUID', output_col]]
        merged = effect_only if merged is None else merged.merge(effect_only, on='CSDDGUID', how='outer')

    total_col = f'S{scenario_num}_Total_Jobs'
    merged[total_col] = merged[effect_cols].sum(axis=1)

    print(f"\nSTEP 4: Saving combined output ({len(effect_cols)} effects + Total, "
          f"single geometry)...")
    output_name = f'scenario{scenario_num}_csd'
    save_outputs(merged, CSD_SHAPEFILE, output_name)
    return merged


if __name__ == '__main__':
    run_scenario_all_effects(6, {
        'DIR': TRAIL_CSD_CSV,          # tariff-relevant NAICS only -- see note in chat re: whether this should change
        'INDIR': TRAIL_CSD_FULL_CSV,   # full economy
        'INDCD': TRAIL_CSD_FULL_CSV,   # full economy
    })
